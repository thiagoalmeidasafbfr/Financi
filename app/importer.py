"""Importa a planilha "Fluxo de Caixa Pessoal" para o banco.

Cada aba "Fluxo Mensal - <nome>" vira um cenário. Para cada linha:
  * linhas manuais (despesas, dívidas avulsas, receita livre...) → valores mensais
    lidos do cache da planilha (resultado das fórmulas);
  * linhas automáticas (salário, 13º, férias, IR, bônus, consignados, quitação) →
    premissas do cenário; células digitadas como constante viram override do mês.
    Se a aba não tiver as premissas (versões antigas), os valores do cache são
    importados como override, preservando o resultado.

Contratos de consignado vêm das abas CLT1/CLT2 e os meses de quitação de cada
cenário vêm da aba Quitações.

Uso:  python -m app.importer caminho/da/planilha.xlsx [--only BR "BR > MEI"]
"""
from __future__ import annotations

import argparse
import re
import unicodedata
from datetime import date, datetime

import openpyxl
from sqlalchemy import select
from sqlalchemy.orm import Session

from . import engine as eng
from .models import (
    Arrear,
    Category,
    InflationRate,
    LoanContract,
    LoanPayoff,
    MonthSetting,
    PlanEntry,
    SalaryChange,
    Scenario,
)

FLOW_PREFIX = "Fluxo Mensal - "

SECTIONS = {
    "RECEITAS": "receita",
    "DESPESAS FIXAS": "despesa_fixa",
    "EMPRÉSTIMOS E DÍVIDAS": "divida",
    "DESPESAS VARIÁVEIS": "despesa_variavel",
    "INVESTIMENTOS E RESERVA": "investimento",
}

# prefixo do rótulo → tipo de linha automática
AUTO_KINDS = [
    ("salário líquido", "salario"),
    ("13º salário", "decimo_terceiro"),
    ("férias (1/3", "ferias"),
    ("restituição ir", "restituicao_ir"),
    ("bônus", "bonus"),
    ("consignados", "consignados"),
    ("quitação antecipada", "quitacao"),
]

# linhas da seção de investimentos que são parâmetros, não categorias
INVEST_PARAM_ROWS = ("% do saldo", "saldo mínimo", "investimento automático")

CONSTANT_FORMULA = re.compile(r"^=[\d.,+\-*/()\s%]+$")


def _norm(text) -> str:
    return unicodedata.normalize("NFC", str(text or "")).strip()


def _month(v) -> date | None:
    if isinstance(v, datetime):
        return date(v.year, v.month, 1)
    if isinstance(v, date):
        return date(v.year, v.month, 1)
    return None


def _num(v) -> float | None:
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def _is_input(formula_value) -> bool:
    """Célula digitada (constante ou fórmula só com números, ex.: =8167-3150)."""
    if formula_value is None:
        return False
    if isinstance(formula_value, str) and formula_value.startswith("="):
        return bool(CONSTANT_FORMULA.match(formula_value))
    return True


class Sheet:
    def __init__(self, ws_f, ws_v):
        self.f, self.v = ws_f, ws_v
        self.labels: dict[str, int] = {}
        for r in range(1, ws_v.max_row + 1):
            label = _norm(ws_v.cell(r, 2).value)
            if label and label not in self.labels:
                self.labels[label] = r

    def row(self, prefix: str) -> int | None:
        p = prefix.lower()
        for label, r in self.labels.items():
            if label.lower().startswith(p):
                return r
        return None

    def params(self, prefix: str, count: int = 1) -> list:
        """Primeiros valores não vazios à direita do rótulo (colunas C..I)."""
        r = self.row(prefix)
        if r is None:
            return [None] * count
        vals = [self.v.cell(r, c).value for c in range(3, 10)]
        vals = [v for v in vals if v not in (None, "")]
        return (vals + [None] * count)[:count]

    def param(self, prefix: str):
        return self.params(prefix)[0]


# --------------------------------------------------------------------------- consignados

def import_loans(db: Session, wb_v, wb_f, report: list[str]) -> dict:
    """Retorna {'contracts': {nome: LoanContract}, 'payoffs': {cenário_quitações: {nome: mês}},
    'sections': [(linha, cenário_quitações)]}."""
    out = {"contracts": {}, "payoffs": {}, "sections": []}
    if "CLT1" not in wb_v.sheetnames:
        report.append("Aba CLT1 não encontrada: nenhum consignado importado.")
        return out
    clt1 = wb_v["CLT1"]
    first_by_name = {}
    if "CLT2" in wb_v.sheetnames:
        clt2 = wb_v["CLT2"]
        for r in range(1, clt2.max_row + 1):
            name, first = _norm(clt2.cell(r, 1).value), _month(clt2.cell(r, 3).value)
            if name and first:
                first_by_name[name] = first

    header = next((r for r in range(1, 20) if _norm(clt1.cell(r, 1).value) == "Contrato"), None)
    r = (header or 4) + 1
    while r <= clt1.max_row:
        name = _norm(clt1.cell(r, 1).value)
        if not name or name.lower() == "total":
            break
        rate, inst, last = _num(clt1.cell(r, 2).value), _num(clt1.cell(r, 3).value), _month(clt1.cell(r, 4).value)
        if inst and last:
            first = first_by_name.get(name) or last
            c = db.scalar(select(LoanContract).where(LoanContract.name == name))
            if c is None:
                c = LoanContract(name=name)
                db.add(c)
            c.creditor = name.split()[0]
            c.monthly_rate, c.installment, c.first_month, c.last_month = rate, inst, first, last
            out["contracts"][name] = c
        r += 1
    db.flush()
    report.append(f"{len(out['contracts'])} contratos de consignado importados.")

    if "Quitações" in wb_v.sheetnames:
        q = wb_v["Quitações"]
        hdr = next((r for r in range(1, 30) if _norm(q.cell(r, 2).value) == "Contrato"), None)
        if hdr:
            for col in range(3, q.max_column + 1):
                h = _norm(q.cell(hdr, col).value)
                m = re.match(r"Mês quitação\s*[–-]\s*(.+)", h)
                if not m:
                    continue
                key = m.group(1).strip()
                pay = {}
                for rr in range(hdr + 1, hdr + 40):
                    name = _norm(q.cell(rr, 2).value)
                    if not name or name.lower() == "total":
                        break
                    month = _month(q.cell(rr, col).value)
                    if month:
                        pay[name] = month
                out["payoffs"][key] = pay
        for rr in range(1, q.max_row + 1):
            m = re.match(r"PARCELAS DESCONTADAS\s*[–-]\s*CENÁRIO\s+(.+)", _norm(q.cell(rr, 2).value), re.I)
            if m:
                out["sections"].append((rr, m.group(1).strip()))
    return out


def _payoff_key_for(formula: str, loans: dict) -> str | None:
    """Descobre qual cenário da aba Quitações a linha de consignados referencia."""
    m = re.search(r"Quitações!\$?C\$?(\d+)", formula or "")
    if not m:
        return None
    row = int(m.group(1))
    key = None
    for start, name in sorted(loans["sections"]):
        if start <= row:
            key = name
    if key is None:
        return None
    for k in loans["payoffs"]:
        if k.lower() == key.lower():
            return k
    return None


# --------------------------------------------------------------------------- cenários

def _category(db: Session, cache: dict, name: str, group: str, kind: str, order: int) -> Category:
    key = (name.lower(), group)
    if key in cache:
        return cache[key]
    cat = Category(name=name, group=group, kind=kind, sort_order=order)
    db.add(cat)
    db.flush()
    cache[key] = cat
    return cat


def _scenario_names(wb_f) -> dict[str, str]:
    """Nome da aba → nome do cenário (usa os rótulos da aba KPIs quando houver)."""
    names = {s: s[len(FLOW_PREFIX):] for s in wb_f.sheetnames if s.startswith(FLOW_PREFIX)}
    if "KPIs" in wb_f.sheetnames:
        k = wb_f["KPIs"]
        for col in range(3, 10):
            label = _norm(k.cell(12, col).value) or _norm(k.cell(4, col).value)
            ref = str(k.cell(14, col).value or "")
            m = re.search(r"'(Fluxo Mensal - [^']+)'", ref)
            if label and m and m.group(1) in names:
                names[m.group(1)] = label
    return names


def import_scenario(db: Session, sheet_name: str, scenario_name: str, wb_f, wb_v, loans: dict,
                    report: list[str]) -> Scenario:
    sh = Sheet(wb_f[sheet_name], wb_v[sheet_name])
    hdr = sh.row("R$")
    if hdr is None:
        raise ValueError(f"{sheet_name}: linha de meses (rótulo 'R$') não encontrada")
    month_cols = [(c, _month(sh.v.cell(hdr, c).value)) for c in range(3, sh.v.max_column + 1)]
    month_cols = [(c, m) for c, m in month_cols if m]
    # meses contíguos a partir do primeiro
    cols: list[tuple[int, date]] = []
    for c, m in month_cols:
        if cols and m != eng.add_months(cols[-1][1], 1):
            break
        cols.append((c, m))
    if not cols:
        raise ValueError(f"{sheet_name}: nenhum mês encontrado")
    # coluna "Em atraso" antes do primeiro mês (valores já vencidos no início do cenário)
    arrear_col = next((c for c in range(3, cols[0][0])
                       if _norm(sh.v.cell(hdr, c).value).lower() == "em atraso"), None)

    old = db.scalar(select(Scenario).where(Scenario.name == scenario_name))
    if old:
        db.delete(old)
        db.flush()

    sc = Scenario(name=scenario_name, description=f"Importado da aba '{sheet_name}'",
                  start_month=cols[0][1], months=len(cols))
    num = lambda v, d=0.0: _num(v) if _num(v) is not None else d  # noqa: E731
    sc.initial_balance = num(sh.param("Saldo inicial em conta"))
    sc.debt_limit_pct = num(sh.param("Limite recomendado de dívidas"), 0.3)
    sc.fgts_initial = num(sh.param("Saldo inicial FGTS"))
    sc.fgts_rate_yearly = num(sh.param("Rendimento FGTS"), 0.03)
    sc.fgts_aliquot = num(sh.param("Alíquota FGTS"), 0.08)
    sc.bonus_salaries = num(sh.param("Bônus anual"))
    sc.bonus_month = sh.param("Mês do bônus")
    sc.ir_amount = num(sh.param("Restituição IR média"))
    sc.ir_month = sh.param("Mês da restituição")
    sc.vacation_month = sh.param("Mês do 1/3 de férias")
    y1, y2 = sh.params("13º salário – primeiro e último ano", 2)
    m1, m2 = sh.params("13º salário – meses", 2)
    sc.thirteenth_first_year, sc.thirteenth_last_year = y1, y2
    sc.thirteenth_month_1, sc.thirteenth_month_2 = m1 or 11, m2 or 12
    sc.inflation_month = sh.param("Mês do reajuste") or 6
    sc.invest_return_monthly = num(sh.param("Rentabilidade padrão"), 0.01)
    sc.initial_invested = num(sh.param("Saldo inicial investido"))
    if sh.row("% do saldo do mês a investir"):
        sc.invest_mode = "saldo_mes"
        sc.invest_pct = num(sh.param("% do saldo do mês a investir"))
    db.add(sc)
    db.flush()

    # inflação
    yr = sh.row("Ano")
    rf, rv = sh.row("Inflação despesas fixas"), sh.row("Inflação despesas variáveis")
    if yr and rf and rv:
        for c in range(3, 10):
            year = _num(sh.v.cell(yr, c).value)
            if year:
                db.add(InflationRate(scenario_id=sc.id, year=int(year), fixed_pct=num(sh.v.cell(rf, c).value),
                                     variable_pct=num(sh.v.cell(rv, c).value)))

    # reajustes de salário
    has_salary_params = False
    r_from, r_amt, r_pct = sh.row("A partir de"), sh.row("Novo salário"), sh.row("ou Reajuste")
    if r_from and r_amt:
        for c in range(3, 10):
            start = _month(sh.v.cell(r_from, c).value)
            amt = _num(sh.v.cell(r_amt, c).value)
            pct = _num(sh.v.cell(r_pct, c).value) if r_pct else None
            if start and (amt is not None or pct is not None):
                db.add(SalaryChange(scenario_id=sc.id, start_month=start, amount=amt, pct=pct))
                has_salary_params = True

    # parâmetros mensais
    gross_row = sh.row("Salário bruto do mês")
    deposit_row = sh.row("(+) Depósito do mês")
    monthly = {
        "invest_pct": sh.row("% do saldo excedente a investir"),
        "min_balance": sh.row("Saldo mínimo a manter em conta"),
        "invest_return": sh.row("Rentabilidade do mês"),
        "manual_redemption": sh.row("(–) Resgate manual"),
    }
    carried = {"invest_pct": sc.invest_pct, "min_balance": 0.0, "invest_return": sc.invest_return_monthly,
               "gross_salary": None}
    for c, m in cols:
        ms = {}
        for key, r in monthly.items():
            if r is None:
                continue
            v = _num(sh.v.cell(r, c).value)
            if key == "manual_redemption":
                if v:
                    ms[key] = v
            elif v is not None and v != carried[key]:
                ms[key] = carried[key] = v
        gross = None
        if gross_row:
            gross = _num(sh.v.cell(gross_row, c).value)
        elif deposit_row:
            dep = _num(sh.v.cell(deposit_row, c).value)
            gross = dep / sc.fgts_aliquot if dep is not None and sc.fgts_aliquot else None
        if gross is not None and gross != carried["gross_salary"]:
            if carried["gross_salary"] is None:
                sc.gross_salary = gross
            else:
                ms["gross_salary"] = gross
            carried["gross_salary"] = gross
        if ms:
            db.add(MonthSetting(scenario_id=sc.id, month=m, **ms))

    # linhas de detalhamento
    cache = {(c.name.lower(), c.group): c for c in db.scalars(select(Category))}
    section, imported_rows, arrears_imported = None, 0, 0
    payoff_key = None
    for r in range(sh.row("DETALHAMENTO") or 1, sh.v.max_row + 1):
        label = _norm(sh.v.cell(r, 2).value)
        if label.upper() in SECTIONS:
            section = SECTIONS[label.upper()]
            continue
        if section is None or not label:
            continue
        if label.lower().startswith("total"):
            section = None
            continue
        if section == "investimento" and label.lower().startswith(INVEST_PARAM_ROWS):
            continue
        name = re.sub(r"\s*\(editável\)\s*$", "", label, flags=re.I)
        kind = "manual"
        if section in ("receita", "divida"):
            kind = next((k for p, k in AUTO_KINDS if name.lower().startswith(p)), "manual")
        if kind != "manual":
            # linhas automáticas usam a categoria única daquele tipo
            existing = next((c for c in cache.values() if c.kind == kind), None)
            name = existing.name if existing else name
        cat = _category(db, cache, name, section, kind, 1000 + r)
        if arrear_col:
            overdue = _num(sh.v.cell(r, arrear_col).value)
            if overdue:
                _upsert_arrear(db, description=name, creditor=None, owner=None, amount=overdue,
                               category_id=cat.id, due_date=eng.add_months(cols[0][1], -1),
                               notes=f"Importado da coluna 'Em atraso' da aba '{sheet_name}'")
                arrears_imported += 1

        # a linha automática pode ser calculada pelo motor?
        computable = False
        if kind == "salario":
            computable = has_salary_params
        elif kind in ("decimo_terceiro", "ferias", "restituicao_ir", "bonus"):
            computable = has_salary_params
        elif kind in ("consignados", "quitacao"):
            formula = str(sh.f.cell(r, cols[0][0]).value or "")
            key = _payoff_key_for(formula, loans)
            computable = key is not None or (not loans["sections"] and "Quitações" not in formula)
            if key is not None:
                payoff_key = key

        for c, m in cols:
            cached = _num(sh.v.cell(r, c).value)
            if kind == "manual" or not computable:
                value = cached if cached else None
            else:
                raw = sh.f.cell(r, c).value
                if raw is None:  # fórmula apagada na planilha = valor zerado naquele mês
                    value = 0.0
                else:
                    value = (cached or 0.0) if _is_input(raw) else None
            if value is not None:
                db.add(PlanEntry(scenario_id=sc.id, category_id=cat.id, month=m, amount=value))
        imported_rows += 1

    if payoff_key:
        for name, month in loans["payoffs"].get(payoff_key, {}).items():
            contract = loans["contracts"].get(name)
            if contract:
                db.add(LoanPayoff(scenario_id=sc.id, contract_id=contract.id, payoff_month=month))

    db.flush()
    report.append(f"Cenário '{scenario_name}' ({sheet_name}): {len(cols)} meses a partir de "
                  f"{cols[0][1]:%m/%Y}, {imported_rows} linhas"
                  + (f", quitações do cenário '{payoff_key}' da aba Quitações" if payoff_key else "")
                  + (f", {arrears_imported} valores em atraso" if arrears_imported else ""))
    sc._validation = (sh, cols)  # usado por validate()
    return sc


def _upsert_arrear(db: Session, description: str, creditor: str | None, owner: str | None, **fields) -> Arrear:
    """Reimportar atualiza o mesmo atraso em vez de duplicar (status e pagamento são preservados)."""
    a = db.scalar(select(Arrear).where(Arrear.description == description,
                                       Arrear.creditor.is_(None) if creditor is None else Arrear.creditor == creditor,
                                       Arrear.owner.is_(None) if owner is None else Arrear.owner == owner))
    if a is None:
        a = Arrear(description=description, creditor=creditor, owner=owner, status="aberto")
        db.add(a)
    for k, v in fields.items():
        setattr(a, k, v)
    db.flush()
    return a


def import_debts_sheet(db: Session, wb_v, report: list[str]) -> None:
    """Aba "Dívidas": linhas no formato  Dívidas <responsável> | credor | valor."""
    if "Dívidas" not in wb_v.sheetnames:
        return
    ws, count, total = wb_v["Dívidas"], 0, 0.0
    for row in ws.iter_rows():
        cells = [c for c in row if c.value not in (None, "")]
        for i, c in enumerate(cells[:-2]):
            m = re.match(r"D[íi]vidas?\s+(.+)", _norm(c.value), re.I)
            if not m:
                continue
            creditor, amount = _norm(cells[i + 1].value), _num(cells[i + 2].value)
            if creditor and amount:
                _upsert_arrear(db, description=creditor, creditor=creditor, owner=m.group(1).strip(),
                               amount=amount, notes="Importado da aba 'Dívidas'")
                count, total = count + 1, total + amount
            break
    if count:
        report.append(f"{count} dívidas em atraso importadas da aba 'Dívidas' (R$ {total:,.2f}).")


def validate(db: Session, sc: Scenario) -> str:
    """Compara a projeção do app com os totais calculados pela planilha."""
    from .services import project_scenario

    sh, cols = sc._validation
    db.refresh(sc)
    p = project_scenario(db, sc)
    checks = [("Receitas totais", "receitas"), ("Total de despesas", "total_despesas"),
              ("Saldo acumulado em conta", "saldo_conta"), ("Saldo total (conta + investimentos)", "saldo_total"),
              ("Saldo FGTS acumulado", "fgts_saldo")]
    ok, parts = [], []
    for label, key in checks:
        r = sh.row(label)
        if r is None:
            continue
        diffs = [((_num(sh.v.cell(r, c).value) or 0.0) - p["series"][key][i], m) for i, (c, m) in enumerate(cols)]
        bad = [(d, m) for d, m in diffs if abs(d) > 0.01]
        if bad:
            worst = max(abs(d) for d, _ in bad)
            parts.append(f"{label}: difere em {len(bad)} mês(es) a partir de {bad[0][1]:%m/%Y} "
                         f"(máx R$ {worst:,.2f})")
        else:
            ok.append(label)
    status = "OK" if not parts else "DIVERGÊNCIAS"
    msg = f"  validação '{sc.name}': {status} — confere: {', '.join(ok) or '-'}"
    return msg + "".join(f"\n    ! {x}" for x in parts)


def import_workbook(db: Session, path, only: list[str] | None = None) -> list[str]:
    wb_f = openpyxl.load_workbook(path)
    wb_v = openpyxl.load_workbook(path, data_only=True)
    report: list[str] = []
    loans = import_loans(db, wb_v, wb_f, report)
    import_debts_sheet(db, wb_v, report)
    names = _scenario_names(wb_f)
    imported = []
    for sheet, name in names.items():
        if only and name not in only and sheet not in only:
            continue
        try:
            imported.append(import_scenario(db, sheet, name, wb_f, wb_v, loans, report))
        except ValueError as exc:
            report.append(f"Aba '{sheet}' ignorada: {exc}")
    if imported and not db.scalar(select(Scenario.id).where(Scenario.is_baseline.is_(True))):
        base = next((s for s in imported if s.name == "BR"), imported[0])
        base.is_baseline = True
    db.commit()
    for sc in imported:
        report.append(validate(db, sc))
    return report


def main() -> None:
    from .db import Base, SessionLocal, engine
    from .services import seed_categories

    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx")
    ap.add_argument("--only", nargs="*", help="importar apenas estes cenários/abas")
    args = ap.parse_args()
    Base.metadata.create_all(engine)
    with SessionLocal() as db:
        seed_categories(db)
        for line in import_workbook(db, args.xlsx, args.only):
            print(line)


if __name__ == "__main__":
    main()
